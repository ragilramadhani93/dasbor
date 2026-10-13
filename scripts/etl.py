#!/usr/bin/env python3
"""ETL: Google Sheet 'KACC Monitoring' -> data/dashboard_data.json"""
import openpyxl, json, re, os, datetime
from collections import Counter, defaultdict

SRC = os.environ.get('SRC', 'sheet.xlsx')
OUT = os.environ.get('OUT', 'data/dashboard_data.json')

MENU_PREFIX = re.compile(r'^(\(RB\)|\(TR\)|\(JK\)|\(SBY\)|ASEAN|TIKTOK|PLAZA BANK INDEX)\s*[-\u2013]?\s*', re.I)

def norm_menu(name):
    if name is None: return ''
    return MENU_PREFIX.sub('', str(name).strip()).strip()

def norm_outlet(name):
    if name is None: return ''
    s = re.sub(r'^KACC\s*[-\u2013]?\s*', '', str(name).strip(), flags=re.I)
    return re.sub(r'\s+', ' ', s).strip()

BAKERY_KEYS = ['PUFF','CROISSANT','PIE','DANISH','BROWNIES','BANANA CAKE','RED VELVET','CINAMON','CINNAMON','MINI BUTTER','APPLE PIE','FRANKIE','PUKIS','WINGKO','LOVY','KASUREN','COOKIES','CAKE']

def category(name):
    n = name.upper()
    if 'PACKAGING FEE' in n: return 'Biaya Kemasan'
    if any(k in n for k in BAKERY_KEYS): return 'Roti, Pastry & Snack'
    if 'CUP 17 AGUSTUS' in n or 'MUG' in n or 'COFFEE PRESS' in n: return 'Merchandise & Promo'
    if 'RTD' in n or 'AIR MINERAL' in n or 'NATA D' in n: return 'RTD & Air Mineral'
    if 'MATCHA' in n: return 'Matcha'
    if any(k in n for k in ['AMERICANO','DRIP','ESPRESSO','AFFOGATO','COLD BREW','BLACK COFFEE','KOPI HITAM','BREW']): return 'Kopi Hitam & Espresso'
    if any(k in n for k in ['LATTE','CAPPUCCINO','COFFEE','KOPI','PISTACHIO','JAVA']): return 'Kopi Susu & Latte'
    if any(k in n for k in ['TEA','THAI','EARL GREY','CHOCO','SMOOTHIES','MELON','ORANGE','COCOPANDAN','PINEAPPLE','BERRY','BARLEY','SUNSET','PURPLE','GOLDEN','CREAMY','FLOAT','ICED','BASE ']): return 'Minuman Non-Kopi'
    return 'Lainnya'

def to_f(v):
    try: return float(v)
    except (TypeError, ValueError): return 0.0

def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    api_txt = None
    for ws in wb.worksheets:
        if ws.title == 'API_Data':
            for row in ws.iter_rows(values_only=True):
                if row and row[0]:
                    api_txt = str(row[0]); break
    api = json.loads(api_txt[api_txt.index('{'):]) if api_txt else {}

    # Delivery Schedule -> master outlet/area
    ws = wb['Delivery Schedule']; it = ws.iter_rows(values_only=True); next(it)
    delivery = []; outlet_area = {}
    for row in it:
        if row is None or row[1] is None: continue
        area = str(row[0]).strip() if row[0] else ''
        outlet = str(row[1]).strip()
        days_raw = str(row[2]).strip() if len(row) > 2 and row[2] else ''
        days = [d.strip() for d in re.split(r'[,;]', days_raw) if d.strip()]
        delivery.append({'area': area, 'outlet': outlet, 'daysRaw': days_raw, 'days': days})
        outlet_area[outlet.upper()] = area
    for o in api.get('outlets', []):
        outlet_area.setdefault(o['outlet'].upper(), o['area'])
    outlet_area.setdefault('OBC UNPAR', 'Jabar')
    outlet_area.setdefault('RS PANTIWILASA', 'Jateng')

    months = ['Jul 2026', 'Aug 2026', 'Sep 2026', 'Oct 2026']
    monthly = []; daily_series = {}
    outlet_month = defaultdict(lambda: {'area': '', 'rev': {}})
    area_month = defaultdict(lambda: defaultdict(float))
    menu_total = Counter(); menu_qty_total = Counter()
    cat_total = Counter(); cat_qty_total = Counter()
    month_menu = {}; payment_month = {}
    hourly_rev = Counter(); hourly_txn = defaultdict(set)
    dow_rev = Counter(); dow_txn = defaultdict(set)
    unmatched_outlets = set()
    DOW = ['Senin','Selasa','Rabu','Kamis','Jumat','Sabtu','Minggu']

    for m in months:
        ws = wb[m]; it = ws.iter_rows(values_only=True); hdr = next(it)
        idx = {h: i for i, h in enumerate(hdr)}
        c_out=idx['Outlet Name']; c_open=idx['Opened Date']; c_txn=idx['Transaction Number']
        c_menu=idx['Item Name']; c_qty=idx['Qty']; c_grand=idx['Grand Total Final']
        c_pay=idx.get('Payment Method')
        day=defaultdict(float); day_txn=defaultdict(set)
        out_rev=Counter(); out_qty=Counter(); out_txn=defaultdict(set)
        area_rev=Counter(); area_txn=defaultdict(set)
        m_rev=Counter(); m_qty=Counter(); pay_rev=Counter(); txn_all=set()
        rev=qty=0.0
        for row in it:
            if row is None or row[c_txn] is None: continue
            outlet_raw = norm_outlet(row[c_out]); key = outlet_raw.upper()
            area = outlet_area.get(key)
            if area is None:
                unmatched_outlets.add(outlet_raw); area = 'Lainnya'
            dt_raw = str(row[c_open]); dstr = dt_raw[:10]; txn = str(row[c_txn])
            menu = norm_menu(row[c_menu]); q = to_f(row[c_qty]); g = to_f(row[c_grand])
            rev += g; qty += q
            day[dstr] += g; day_txn[dstr].add(txn)
            out_rev[outlet_raw] += g; out_qty[outlet_raw] += q; out_txn[outlet_raw].add(txn)
            area_rev[area] += g; area_txn[area].add(txn)
            m_rev[menu] += g; m_qty[menu] += q
            menu_total[menu] += g; menu_qty_total[menu] += q
            c = category(menu); cat_total[c] += g; cat_qty_total[c] += q
            if c_pay is not None:
                pay = str(row[c_pay]).strip().replace('\n', ' ').split('\u2022')[-1].strip()
                pay_rev[pay] += g
            txn_all.add(txn)
            try:
                hh = int(dt_raw[11:13]); hourly_rev[hh] += g; hourly_txn[hh].add(txn)
            except Exception: pass
            try:
                dd = datetime.date(int(dstr[0:4]), int(dstr[5:7]), int(dstr[8:10]))
                dw = dd.weekday(); dow_rev[dw] += g; dow_txn[dw].add(txn)
            except Exception: pass
        monthly.append({'month': m, 'revenue': rev, 'qty': qty, 'transactions': len(txn_all),
                        'avgBasket': (rev/len(txn_all)) if txn_all else 0, 'outlets': len(out_rev)})
        daily_series[m] = [{'date': d, 'revenue': day[d], 'transactions': len(day_txn[d])} for d in sorted(day)]
        for o, r in out_rev.items():
            outlet_month[o]['rev'][m] = r
            outlet_month[o]['area'] = outlet_area.get(o.upper(), 'Lainnya')
        for a, r in area_rev.items(): area_month[a][m] = r
        month_menu[m] = [{'menu': k, 'revenue': v, 'qty': m_qty[k]} for k, v in m_rev.most_common(25)]
        payment_month[m] = [{'method': k, 'revenue': v} for k, v in pay_rev.most_common()]

    outlet_month_table = []
    for o, d in outlet_month.items():
        row_d = {'outlet': o, 'area': d['area']}; row_d.update(d['rev'])
        outlet_month_table.append(row_d)
    outlet_month_table.sort(key=lambda r: -sum(v for k, v in r.items() if k not in ('outlet', 'area')))
    area_month_table = []
    for a, d in area_month.items():
        row_d = {'area': a}; row_d.update(d); area_month_table.append(row_d)

    # Stock balance
    ws = wb['Stock_Balance']; it = ws.iter_rows(values_only=True); next(it)
    stock_rows = []; stock_by_outlet = defaultdict(lambda: {'neg': 0, 'zero': 0, 'pos': 0, 'min': 0})
    for row in it:
        if row is None or row[0] is None: continue
        outlet = str(row[0]).strip(); menu = str(row[1]).strip(); v = to_f(row[2])
        lu = str(row[3]).strip() if len(row) > 3 and row[3] else ''
        stock_rows.append({'outlet': outlet, 'menu': menu, 'stock': v,
                           'area': outlet_area.get(outlet.upper(), 'Lainnya'), 'lastUpdated': lu})
        s = stock_by_outlet[outlet]
        if v < 0: s['neg'] += 1
        elif v == 0: s['zero'] += 1
        else: s['pos'] += 1
        s['min'] = min(s['min'], v)

    # BOM
    ws = wb['BOM Master']; it = ws.iter_rows(values_only=True); next(it)
    bom_ing = Counter(); bom = []
    for row in it:
        if row is None or not row[0]: continue
        items = []
        for b in range(2, 26, 3):
            if b >= len(row): break
            if row[b]:
                items.append({'bahan': str(row[b]), 'qty': to_f(row[b+1]), 'satuan': str(row[b+2])})
                bom_ing[str(row[b])] += 1
        bom.append({'code': str(row[0]).replace('.0', ''), 'menu': str(row[1]), 'items': items})

    # Ingredient usage
    ws = wb['Ingredient Usage']; it = ws.iter_rows(values_only=True); next(it)
    iu_by_ing = Counter(); iu_ing_unit = {}; iu_by_area = Counter()
    iu_outlets = set(); iu_periods = set(); iu_rows = 0; iu_placeholder = 0
    for row in it:
        if row is None or row[0] is None: continue
        iu_rows += 1; bahan = str(row[4])
        if 'belum diisi' in bahan:
            iu_placeholder += 1; continue
        val = to_f(row[6]); iu_by_ing[bahan] += val; iu_ing_unit[bahan] = str(row[5])
        iu_by_area[str(row[0])] += val; iu_outlets.add(str(row[1])); iu_periods.add(str(row[2])[:10])

    def read_log(name):
        ws = wb[name]; it = ws.iter_rows(values_only=True)
        hdr = [h for h in next(it) if h is not None]; rows = []
        for row in it:
            if row is None or row[0] is None: continue
            rows.append([('' if c is None else str(c)) for c in row])
        return {'headers': [str(h) for h in hdr], 'rows': rows}
    logs = {name: read_log(name) for name in ['GR_Log', 'ReturWaste_Log', 'Opname_Log']}

    result = {
        'meta': {'generatedAt': datetime.datetime.utcnow().isoformat() + 'Z',
                 'source': 'Google Sheet KACC Monitoring',
                 'reportDate': api.get('dateStr', '12 Oct 2026'),
                 'areas': api.get('areas', []), 'months': months, 'currency': 'IDR'},
        'sales': {'daily': api.get('sales', {}).get('daily', []),
                  'mtd': api.get('sales', {}).get('mtd', []),
                  'snack': api.get('snack', {})},
        'trend': {'monthly': monthly, 'dailySeries': daily_series,
                  'outletMonthly': outlet_month_table, 'areaMonthly': area_month_table,
                  'topMenus': [{'menu': k, 'revenue': v, 'qty': menu_qty_total[k], 'category': category(k)} for k, v in menu_total.most_common(30)],
                  'categoryMix': [{'category': k, 'revenue': v, 'qty': cat_qty_total[k]} for k, v in cat_total.most_common()],
                  'paymentMix': payment_month,
                  'hourly': [{'hour': h, 'revenue': hourly_rev[h], 'transactions': len(hourly_txn[h])} for h in range(24)],
                  'dow': [{'dow': DOW[d], 'revenue': dow_rev[d], 'transactions': len(dow_txn[d])} for d in range(7)],
                  'menuByMonth': month_menu},
        'inventory': {'stock': {'rows': stock_rows,
                                'byOutlet': [{'outlet': o, **v} for o, v in sorted(stock_by_outlet.items(), key=lambda x: x[1]['neg'], reverse=True)]},
                      'bom': {'count': len(bom), 'ingredients': [{'bahan': k, 'count': v} for k, v in bom_ing.most_common()], 'items': bom},
                      'ingredientUsage': {'rows': iu_rows, 'placeholderRows': iu_placeholder,
                                          'ingredients': [{'bahan': k, 'total': v, 'satuan': iu_ing_unit[k]} for k, v in iu_by_ing.most_common()],
                                          'byArea': [{'area': k, 'total': v} for k, v in iu_by_area.most_common()],
                                          'outlets': sorted(iu_outlets), 'periods': sorted(iu_periods)},
                      'delivery': delivery},
        'logs': logs,
        'diagnostics': {'unmatchedOutlets': sorted(unmatched_outlets)},
    }
    os.makedirs(os.path.dirname(OUT) or '.', exist_ok=True)
    with open(OUT, 'w') as f:
        json.dump(result, f, ensure_ascii=False, separators=(',', ':'))
    print('Wrote', OUT, os.path.getsize(OUT), 'bytes')
    print('unmatched outlets:', sorted(unmatched_outlets))
    print('monthly:', [(x['month'], round(x['revenue'])) for x in monthly])
    print('categories:', [(k, round(v)) for k, v in cat_total.most_common()])

if __name__ == '__main__':
    main()
