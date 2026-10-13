#!/usr/bin/env python3
"""Bangun Monitoring_Dashboard.xlsx (versi spreadsheet) dari data/dashboard_data.json."""
import json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, PieChart, Reference

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
D = json.load(open(os.path.join(BASE, 'data/dashboard_data.json')))
OUT = os.path.join(BASE, 'Monitoring_Dashboard.xlsx')

HDR = PatternFill('solid', fgColor='0F6F5C')
HDRF = Font(color='FFFFFF', bold=True)
TITLE = Font(bold=True, size=14, color='0B2F2A')
SUB = Font(italic=True, size=10, color='6B7A90')
THIN = Side(style='thin', color='D9E2EC')
BORD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CUR = '#,##0'

def head(ws, row, cols):
    for j, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.fill = HDR; cell.font = HDRF; cell.border = BORD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

def autow(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

def title(ws, t, s=''):
    ws['A1'] = t; ws['A1'].font = TITLE
    if s: ws['A2'] = s; ws['A2'].font = SUB

wb = Workbook()

# ---------- Dashboard (KPI) ----------
ws = wb.active; ws.title = 'Dashboard'
title(ws, 'DASHBOARD MONITORING KACC', 'Per '+str(D['meta'].get('reportDate'))+'  |  sumber: Google Sheet')
daily = D['sales']['daily']; mtd = D['sales']['mtd']
dRev = sum(x['revenue'] for x in daily); dTar = sum(x['target'] for x in daily)
mRev = sum(x['revenue'] for x in mtd); mTar = sum(x['target'] for x in mtd)
snat = D['sales']['snack'].get('byMenuNational', [])
rows = [
    ('Omset Harian', dRev), ('Target Harian', dTar), ('Pencapaian Harian', dRev/dTar if dTar else 0),
    ('Omset MTD', mRev), ('Target MTD', mTar), ('Pencapaian MTD', mRev/mTar if mTar else 0),
    ('Omset Snack Harian', sum(x['dRev'] for x in snat)), ('Omset Snack MTD', sum(x['mRev'] for x in snat)),
    ('Jumlah Outlet', len(daily)), ('Outlet Aktif (omset>0)', sum(1 for x in daily if x['revenue'] > 0)),
]
head(ws, 4, ['Indikator', 'Nilai'])
r = 5
for k, v in rows:
    ws.cell(row=r, column=1, value=k).border = BORD
    c = ws.cell(row=r, column=2, value=v); c.border = BORD
    c.number_format = '0.0%' if 'Pencapaian' in k else CUR
    r += 1
autow(ws, [28, 18])
# data area untuk chart
ws.cell(row=4, column=4, value='Area').font = HDRF
ws.cell(row=4, column=4).fill = HDR
ws.cell(row=4, column=5, value='Realisasi Harian').font = HDRF; ws.cell(row=4, column=5).fill = HDR
ws.cell(row=4, column=6, value='Target Harian').font = HDRF; ws.cell(row=4, column=6).fill = HDR
ws.cell(row=4, column=7, value='Omset MTD').font = HDRF; ws.cell(row=4, column=7).fill = HDR
AREAS = D['meta']['areas']
rr = 5
for a in AREAS:
    ws.cell(row=rr, column=4, value=a)
    ws.cell(row=rr, column=5, value=sum(x['revenue'] for x in daily if x['area'] == a)).number_format = CUR
    ws.cell(row=rr, column=6, value=sum(x['target'] for x in daily if x['area'] == a)).number_format = CUR
    ws.cell(row=rr, column=7, value=sum(x['revenue'] for x in mtd if x['area'] == a)).number_format = CUR
    rr += 1
ch = BarChart(); ch.title = 'Realisasi vs Target Harian per Area'; ch.height = 8; ch.width = 14
ch.add_data(Reference(ws, min_col=5, max_col=6, min_row=4, max_row=rr-1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=4, min_row=5, max_row=rr-1))
ws.add_chart(ch, 'I4')
pie = PieChart(); pie.title = 'Kontribusi Omset MTD per Area'; pie.height = 8; pie.width = 14
pie.add_data(Reference(ws, min_col=7, min_row=4, max_row=rr-1), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=4, min_row=5, max_row=rr-1))
ws.add_chart(pie, 'I22')
print('Dashboard ok')

# ---------- Performa Outlet ----------
ws = wb.create_sheet('Outlet')
title(ws, 'PERFORMA OUTLET', 'Harian & MTD per '+str(D['meta'].get('reportDate')))
head(ws, 4, ['Area', 'Outlet', 'Omset Harian', 'Target Harian', '% Harian', 'Omset MTD', 'Target MTD', '% MTD', 'Shortfall Harian'])
md = {x['outlet']: x for x in mtd}
rr = 5
for x in sorted(daily, key=lambda z: -md.get(z['outlet'], {}).get('revenue', 0)):
    m = md.get(x['outlet'], {})
    vals = [x['area'], x['outlet'], x['revenue'], x['target'], x['pct'], m.get('revenue', 0), m.get('target', 0), m.get('pct', 0), x['shortfall']]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rr, column=j, value=v); c.border = BORD
        if j in (3, 4, 6, 7, 9): c.number_format = CUR
        if j in (5, 8): c.number_format = '0.0%'
    rr += 1
autow(ws, [14, 26, 16, 16, 10, 16, 16, 10, 16])
ws.freeze_panes = 'A5'
ch = BarChart(); ch.type = 'bar'; ch.title = 'Omset MTD per Outlet'; ch.height = 18; ch.width = 16
ch.add_data(Reference(ws, min_col=6, min_row=4, max_row=min(rr-1, 4+20)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=2, min_row=5, max_row=min(rr-1, 4+20)))
ws.add_chart(ch, 'K4')

# ---------- Tren ----------
ws = wb.create_sheet('Tren')
title(ws, 'TREN PENJUALAN', 'Bulanan & per outlet')
head(ws, 4, ['Bulan', 'Omset', 'Qty', 'Transaksi', 'Rata2/Transaksi', 'Outlet'])
rr = 5
for x in D['trend']['monthly']:
    for j, v in enumerate([x['month'], x['revenue'], x['qty'], x['transactions'], round(x['avgBasket']), x['outlets']], 1):
        c = ws.cell(row=rr, column=j, value=v); c.border = BORD
        if j in (2, 5): c.number_format = CUR
    rr += 1
autow(ws, [14, 18, 12, 12, 16, 10])
ch = LineChart(); ch.title = 'Tren Omset Bulanan'; ch.height = 8; ch.width = 14
ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=rr-1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rr-1))
ws.add_chart(ch, 'H4')
# outlet x month
base = rr + 2
ws.cell(row=base, column=1, value='Omset per Outlet per Bulan').font = TITLE
head(ws, base+1, ['Area', 'Outlet'] + D['meta']['months'] + ['Total'])
rr2 = base + 2
for o in D['trend']['outletMonthly']:
    tot = sum(o.get(m, 0) for m in D['meta']['months'])
    vals = [o['area'], o['outlet']] + [o.get(m, 0) for m in D['meta']['months']] + [tot]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=rr2, column=j, value=v); c.border = BORD
        if j >= 3: c.number_format = CUR
    rr2 += 1
ws.freeze_panes = 'A6'
print('Outlet/Tren ok')

# ---------- Produk ----------
ws = wb.create_sheet('Produk')
title(ws, 'PRODUK, KATEGORI & PEMBAYARAN', 'Seluruh periode')
head(ws, 4, ['Kategori', 'Omset', 'Qty'])
rr = 5
for c in D['trend']['categoryMix']:
    for j, v in enumerate([c['category'], c['revenue'], c['qty']], 1):
        cell = ws.cell(row=rr, column=j, value=v); cell.border = BORD
        if j == 2: cell.number_format = CUR
    rr += 1
autow(ws, [28, 18, 12])
pie = PieChart(); pie.title = 'Bauran Kategori (Omset)'; pie.height = 9; pie.width = 14
pie.add_data(Reference(ws, min_col=2, min_row=4, max_row=rr-1), titles_from_data=True)
pie.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rr-1))
ws.add_chart(pie, 'E4')
base = rr + 2
ws.cell(row=base, column=1, value='Top 20 Produk').font = TITLE
head(ws, base+1, ['Produk', 'Kategori', 'Omset', 'Qty'])
r2 = base + 2
for m in D['trend']['topMenus'][:20]:
    for j, v in enumerate([m['menu'], m['category'], m['revenue'], m['qty']], 1):
        cell = ws.cell(row=r2, column=j, value=v); cell.border = BORD
        if j == 3: cell.number_format = CUR
    r2 += 1
ch = BarChart(); ch.type = 'bar'; ch.title = 'Top Produk (Omset)'; ch.height = 12; ch.width = 14
ch.add_data(Reference(ws, min_col=3, min_row=base+1, max_row=r2-1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=base+2, max_row=r2-1))
ws.add_chart(ch, 'E20')
base2 = r2 + 2
ws.cell(row=base2, column=1, value='Metode Pembayaran (bulan terakhir)').font = TITLE
head(ws, base2+1, ['Metode', 'Omset'])
mth = D['meta']['months'][-1]
r3 = base2 + 2
for p in D['trend']['paymentMix'].get(mth, []):
    ws.cell(row=r3, column=1, value=p['method']).border = BORD
    c = ws.cell(row=r3, column=2, value=p['revenue']); c.border = BORD; c.number_format = CUR
    r3 += 1

# ---------- Snack ----------
ws = wb.create_sheet('Snack')
title(ws, 'SNACK / ROTI & PASTRY', 'Nasional & per area')
head(ws, 4, ['Menu', 'Qty Harian', 'Nominal Harian', 'Qty MTD', 'Nominal MTD'])
rr = 5
for m in D['sales']['snack'].get('byMenuNational', []):
    for j, v in enumerate([m['menu'], m['dQty'], m['dRev'], m['mQty'], m['mRev']], 1):
        c = ws.cell(row=rr, column=j, value=v); c.border = BORD
        if j in (3, 5): c.number_format = CUR
    rr += 1
autow(ws, [30, 12, 16, 12, 16])
ch = BarChart(); ch.type = 'bar'; ch.title = 'Snack per Menu (MTD)'; ch.height = 12; ch.width = 14
ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=rr-1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rr-1))
ws.add_chart(ch, 'G4')
base = rr + 2
ws.cell(row=base, column=1, value='Detail per Outlet').font = TITLE
head(ws, base+1, ['Area', 'Outlet', 'Menu', 'Qty Harian', 'Nominal Harian', 'Qty MTD', 'Nominal MTD'])
r2 = base + 2
for d in D['sales']['snack'].get('detail', []):
    for j, v in enumerate([d['area'], d['outlet'], d['menu'], d['dQty'], d['dRev'], d['mQty'], d['mRev']], 1):
        c = ws.cell(row=r2, column=j, value=v); c.border = BORD
        if j in (5, 7): c.number_format = CUR
    r2 += 1
print('Produk/Snack ok')

# ---------- Inventori ----------
ws = wb.create_sheet('Inventori')
title(ws, 'INVENTORI & STOK', 'Saldo stok & anomali negatif')
head(ws, 4, ['Outlet', 'SKU Negatif', 'SKU Nol', 'SKU Positif', 'Stok Terendah'])
rr = 5
for o in D['inventory']['stock']['byOutlet']:
    for j, v in enumerate([o['outlet'], o['neg'], o['zero'], o['pos'], o['min']], 1):
        c = ws.cell(row=rr, column=j, value=v); c.border = BORD
    rr += 1
autow(ws, [26, 12, 10, 12, 14])
ch = BarChart(); ch.type = 'bar'; ch.title = 'SKU Negatif per Outlet'; ch.height = 12; ch.width = 14
ch.add_data(Reference(ws, min_col=2, min_row=4, max_row=rr-1), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=rr-1))
ws.add_chart(ch, 'G4')
base = rr + 2
ws.cell(row=base, column=1, value='Detail Stok Negatif').font = TITLE
head(ws, base+1, ['Area', 'Outlet', 'Menu', 'Stok', 'Update Terakhir'])
r2 = base + 2
for s in sorted([x for x in D['inventory']['stock']['rows'] if x['stock'] < 0], key=lambda z: z['stock']):
    for j, v in enumerate([s['area'], s['outlet'], s['menu'], s['stock'], s['lastUpdated']], 1):
        ws.cell(row=r2, column=j, value=v).border = BORD
    r2 += 1

# ---------- Bahan Baku ----------
ws = wb.create_sheet('Bahan Baku')
title(ws, 'BAHAN BAKU & BOM', 'Pemakaian & resep')
head(ws, 4, ['Bahan Baku', 'Satuan', 'Total Pemakaian'])
rr = 5
for i in D['inventory']['ingredientUsage']['ingredients']:
    for j, v in enumerate([i['bahan'], i['satuan'], i['total']], 1):
        c = ws.cell(row=rr, column=j, value=v); c.border = BORD
        if j == 3: c.number_format = CUR
    rr += 1
autow(ws, [40, 10, 18])
ch = BarChart(); ch.type = 'bar'; ch.title = 'Top Pemakaian Bahan Baku'; ch.height = 12; ch.width = 14
ch.add_data(Reference(ws, min_col=3, min_row=4, max_row=min(rr-1, 4+15)), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=min(rr-1, 4+15)))
ws.add_chart(ch, 'E4')
base = rr + 2
ws.cell(row=base, column=1, value='Daftar Resep (BOM Master)').font = TITLE
head(ws, base+1, ['Kode', 'Menu', 'Komposisi'])
r2 = base + 2
for b in D['inventory']['bom']['items']:
    comp = ', '.join(f"{i['bahan']} ({i['qty']:g} {i['satuan']})" for i in b['items'])
    for j, v in enumerate([b['code'], b['menu'], comp], 1):
        ws.cell(row=r2, column=j, value=v).border = BORD
    r2 += 1

# ---------- Logistik ----------
ws = wb.create_sheet('Logistik')
title(ws, 'LOGISTIK & LOG OPERASIONAL', 'Jadwal pengiriman & log')
head(ws, 4, ['Area', 'Outlet', 'Hari Kirim'])
rr = 5
for d in D['inventory']['delivery']:
    vals = [d['area'], d['outlet'], ', '.join(d['days']) if d['days'] else 'Belum diisi']
    for j, v in enumerate(vals, 1):
        ws.cell(row=rr, column=j, value=v).border = BORD
    rr += 1
autow(ws, [14, 26, 40])
base = rr + 2
for name, L in D['logs'].items():
    ws.cell(row=base, column=1, value=name).font = TITLE
    head(ws, base+1, L['headers'] if L['headers'] else ['(kosong)'])
    if not L['rows']:
        ws.cell(row=base+2, column=1, value='Belum ada data tercatat').font = SUB
        base += 4
    else:
        r2 = base + 2
        for row in L['rows']:
            for j, v in enumerate(row, 1):
                ws.cell(row=r2, column=j, value=v).border = BORD
            r2 += 1
        base = r2 + 2

wb.save(OUT)
print('Wrote', OUT, os.path.getsize(OUT), 'bytes')
