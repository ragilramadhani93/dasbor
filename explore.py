import openpyxl, json, sys

wb = openpyxl.load_workbook('sheet.xlsx', read_only=True, data_only=True)
out = []
for ws in wb.worksheets:
    info = {'title': ws.title, 'max_row': ws.max_row, 'max_col': ws.max_column}
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows.append(row)
        if i >= 20:
            break
    info['sample'] = rows
    out.append(info)

with open('explore_out.json', 'w') as f:
    json.dump(out, f, default=str, indent=1)

for info in out:
    print('=' * 80)
    print(f"SHEET: {info['title']}  rows={info['max_row']} cols={info['max_col']}")
    for i, r in enumerate(info['sample']):
        cells = [('' if c is None else str(c)) for c in r]
        # trim trailing empties
        while cells and cells[-1] == '':
            cells.pop()
        if cells:
            print(f"  r{i}: {cells}")
