import openpyxl, json
wb = openpyxl.load_workbook('sheet.xlsx', read_only=True, data_only=True)
# Delivery Schedule full
ws = wb['Delivery Schedule']
print('DELIVERY SCHEDULE dims', ws.calculate_dimension(force=True))
n=0
for row in ws.iter_rows(values_only=True):
    cells=[('' if c is None else str(c)) for c in row]
    while cells and cells[-1]=='': cells.pop()
    if not cells: continue
    print(f'  [{n}] {cells}')
    n+=1
    if n>=25: break
print('total-ish n', n)
print()
ws = wb['Summary']
print('SUMMARY full')
n=0
for row in ws.iter_rows(values_only=True):
    cells=[('' if c is None else str(c)) for c in row]
    while cells and cells[-1]=='': cells.pop()
    print(f'  [{n}] {cells}')
    n+=1
    if n>=260: break
