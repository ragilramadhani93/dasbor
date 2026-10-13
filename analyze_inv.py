import openpyxl
from collections import Counter, defaultdict
wb = openpyxl.load_workbook('sheet.xlsx', read_only=True, data_only=True)
# Ingredient usage
ws = wb['Ingredient Usage']
it=ws.iter_rows(values_only=True); hdr=next(it)
print('IU hdr', hdr)
ingr=Counter(); units=defaultdict(set); areas=Counter(); outlets=set(); periods=set(); rows=0; ph=0
for row in it:
    if row is None or row[0] is None: continue
    rows+=1
    b=str(row[4]); u=str(row[5]); 
    if 'belum diisi' in b: ph+=1; continue
    ingr[b]+=float(row[6] or 0); units[b].add(u); areas[row[0]]+=1; outlets.add(row[1]); periods.add((str(row[2])[:10],str(row[3])[:10]))
print('rows',rows,'placeholder',ph,'unique ingredients',len(ingr))
print('periods',sorted(periods)[:10])
print('areas',areas)
print('top ingredients', ingr.most_common(15))
print()
# BOM
ws=wb['BOM Master']
it=ws.iter_rows(values_only=True); hdr=next(it)
print('BOM hdr len', len(hdr))
n=0; ings=Counter(); units2=defaultdict(set); maxb=0
for row in it:
    if row is None or row[0] is None: continue
    n+=1
    for b in range(2,26,3):
        name=row[b]
        if name: ings[str(name)]+=1; units2[str(name)].add(str(row[b+2]))
print('BOM menus',n,'unique bahan',len(ings))
print('top bahan',ings.most_common(10))
print()
# Stock balance
ws=wb['Stock_Balance']
it=ws.iter_rows(values_only=True); hdr=next(it); print('SB hdr',hdr)
n=0; neg=0; zero=0; pos=0; byout=Counter(); negdet=[]
for row in it:
    if row is None or row[0] is None: continue
    n+=1; v=float(row[2] or 0); byout[row[0]]+=1
    if v<0: neg+=1; negdet.append((row[0],row[1],v))
    elif v==0: zero+=1
    else: pos+=1
print('SB rows',n,'neg',neg,'zero',zero,'pos',pos,'outlets',len(byout))
print('neg sample',negdet[:15])
print('outlet counts', dict(byout))
