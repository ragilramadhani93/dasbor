import openpyxl
from collections import Counter, defaultdict
wb = openpyxl.load_workbook('sheet.xlsx', read_only=True, data_only=True)
ws=wb['BOM Master']
it=ws.iter_rows(values_only=True); hdr=next(it)
print('BOM hdr len', len(hdr))
n=0; ings=Counter(); units2=defaultdict(set); permenu=[]
for row in it:
    if row is None or not row[0]: continue
    n+=1; cnt=0
    for b in range(2,26,3):
        if b>=len(row): break
        name=row[b]
        if name: ings[str(name)]+=1; units2[str(name)].add(str(row[b+2])); cnt+=1
    permenu.append(cnt)
print('BOM menus',n,'unique bahan',len(ings),'avg bahan/menu',sum(permenu)/len(permenu))
print('top bahan',ings.most_common(12))
print('units sample', list(units2.items())[:8])
print()
ws=wb['Stock_Balance']
it=ws.iter_rows(values_only=True); hdr=next(it); print('SB hdr',hdr)
n=0; neg=0; zero=0; pos=0; byout=Counter(); negdet=[]; posdet=[]
for row in it:
    if row is None or not row[0]: continue
    n+=1; v=float(row[2] or 0); byout[row[0]]+=1
    if v<0: neg+=1; negdet.append((row[0],row[1],v))
    elif v==0: zero+=1
    else: pos+=1; posdet.append((row[0],row[1],v))
print('SB rows',n,'neg',neg,'zero',zero,'pos',pos,'outlets',len(byout))
print('neg sample',negdet[:20])
print('pos sample',posdet[:20])
print('outlet list', list(byout.items()))
